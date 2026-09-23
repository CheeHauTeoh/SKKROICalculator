# Builds SK_Keong_Pricing_App_Data_Request.xlsx: bilingual fill-in forms whose headers match the
# app importer (canonical_name 中文 -> the importer ignores the Chinese suffix).
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

FONT = 'Arial'
wb = Workbook()
YELLOW = PatternFill('solid', fgColor='FFFF00'); GREY = PatternFill('solid', fgColor='E7E6E6'); GREEN = PatternFill('solid', fgColor='E2EFDA'); BLUE = PatternFill('solid', fgColor='DDEBF7')
thin = Side(style='thin', color='BFBFBF'); BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
H = Font(name=FONT, bold=True, size=10); N = Font(name=FONT, size=10); T = Font(name=FONT, bold=True, size=14); S = Font(name=FONT, size=9, color='595959'); INPUT = Font(name=FONT, size=10, color='0000FF')

def title(ws, zh, en, note_zh, note_en):
    ws['A1'] = f'{zh}  |  {en}'; ws['A1'].font = T
    ws['A2'] = note_zh; ws['A2'].font = S
    ws['A3'] = note_en; ws['A3'].font = S
    ws.row_dimensions[1].height = 22

def header(ws, row, cols):
    """cols: list of (header, width, comment_en, fill_input:boolean)"""
    for i, (h, w, c, inp) in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=h); cell.font = H; cell.fill = YELLOW if inp else GREY; cell.border = BORDER
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        if c: cell.comment = Comment(c, 'Pricing app')
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 42
    ws.freeze_panes = ws.cell(row=row + 1, column=1)

def rows(ws, start, data, ncols, input_cols, n_blank=0):
    r = start
    for rec in data:
        for i in range(1, ncols + 1):
            v = rec[i - 1] if i - 1 < len(rec) else None
            c = ws.cell(row=r, column=i, value=v); c.font = INPUT if i in input_cols else N; c.border = BORDER
            if i in input_cols: c.fill = YELLOW
        r += 1
    for _ in range(n_blank):
        for i in range(1, ncols + 1):
            c = ws.cell(row=r, column=i); c.border = BORDER; c.font = INPUT if i in input_cols else N
            if i in input_cols: c.fill = YELLOW
        r += 1
    return r

def legend(ws, row):
    ws.cell(row=row, column=1, value='图例 Legend').font = H
    ws.cell(row=row + 1, column=1).fill = YELLOW; ws.cell(row=row + 1, column=2, value='黄色 = 请填写 Yellow = please fill in').font = N
    ws.cell(row=row + 2, column=1).fill = GREY; ws.cell(row=row + 2, column=2, value='灰色 = 系统提供，请勿修改 Grey = provided, do not change').font = N
    ws.cell(row=row + 3, column=1).fill = GREEN; ws.cell(row=row + 3, column=2, value='绿色 = 示例行，请覆盖或删除 Green = example row, overwrite or delete').font = N

# ------------------------------------------------------------------ 0 说明
ws = wb.active; ws.title = '0 说明 Instructions'
ws.column_dimensions['A'].width = 34; ws.column_dimensions['B'].width = 70; ws.column_dimensions['C'].width = 26; ws.column_dimensions['D'].width = 16
title(ws, 'SK Keong 价格应用 · 资料收集表', 'SK Keong pricing app · data request',
      '请各负责人填写对应的工作表（黄色格）。填好后把整个文件发回，或把每张表另存为 CSV 后在应用的「导入」页面上传。',
      'Each owner fills in their sheet (yellow cells). Return the whole workbook, or save each sheet as CSV and upload it on the app\'s Import page.')
r = 5
ws.cell(row=r, column=1, value='工作表 Sheet').font = H; ws.cell(row=r, column=2, value='需要什么 What we need').font = H; ws.cell(row=r, column=3, value='谁填 Who').font = H; ws.cell(row=r, column=4, value='优先级 Priority').font = H
for c in range(1, 5): ws.cell(row=r, column=c).fill = GREY; ws.cell(row=r, column=c).border = BORDER
plan = [
    ('1 用户 Users', '每位使用者的登录名、姓名、角色。业务员代码必须与 UBS 一致。\nLogin, name and role for every user. Salesperson code must match UBS.', '老板 Owner / 文员 Clerk', '第1周 Week 1'),
    ('2 单位核对 UOM', '里程碑1：核心SKU的采购单位、销售单位、换算系数、核实成本。这是整个项目的前提。\nMilestone 1: purchase unit, selling unit, factor and verified cost for the core SKUs. Nothing else works without this.', '财务 Finance (Yun Jun) + 采购 Purchasing', '第1–2周 Week 1–2'),
    ('3 价格等级 Tiers', '价格等级的名称，以及「客户类型 × 规模 → 等级」的规则。\nNames of the price tiers, and the rule mapping customer type × size to a tier.', '老板 Owner', '第2周 Week 2'),
    ('4 目标毛利 Margins', '每个品类的目标毛利率和底价折扣。\nTarget margin and floor discount per category.', '老板 Owner + 财务 Finance', '第2周 Week 2'),
    ('5 核心价格 Prices', '核心SKU现有的目录价和底价（每个等级）。如果目前没有底价，只填目录价。\nCurrent list and floor price per tier for the core SKUs. If there is no floor today, fill list only.', '财务 Finance', '第2–3周 Week 2–3'),
    ('6 客户等级 Customers', '大客户 / 特殊客户的等级手动指定，以及客户类型、规模的确认。\nManual tier for key/special customers; confirm customer type and size.', '老板 Owner + 业务员 Sales', '第3周 Week 3'),
    ('7 竞争对手 Competitors', '主要竞争对手名单（记录竞争价时用），品牌货的市场参考来源。\nMain competitors (used when capturing competitor prices); market-reference sources for branded SKUs.', '业务员 Sales + 老板 Owner', '第3周 Week 3'),
    ('8 决策 Decisions', '老板需要拍板的规则（如业务员能否看到成本）。\nRules the owner must decide (e.g. whether salespeople may see cost).', '老板 Owner', '第1周 Week 1'),
]
r += 1
for p in plan:
    for c, v in enumerate(p, 1):
        cell = ws.cell(row=r, column=c, value=v); cell.font = N; cell.border = BORDER; cell.alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[r].height = 48; r += 1
r += 1
ws.cell(row=r, column=1, value='如何填写 How to fill').font = H; r += 1
for line in [
    '1. 只改黄色格。灰色格是系统提供的数据，请勿改动。 Only edit yellow cells. Grey cells are provided; do not change them.',
    '2. 每张表第一行是列名，请勿改动（导入时需要）。 Row 1 of each data table is the column name; keep it exactly (the import needs it).',
    '3. 金额用马币，两位小数，不要加 "RM"。 Money in RM with two decimals, without the "RM" prefix.',
    '4. 绿色示例行请覆盖或删除。 Overwrite or delete the green example rows.',
    '5. 不确定的地方，在「备注」写下问题，不要猜。 If unsure, write the question in the notes column; do not guess.',
    '6. 交回：整个文件发给项目负责人，或在应用「导入 / 更新」页面上传另存为 CSV 的工作表。 Return: send the workbook back, or upload each sheet saved as CSV on the app\'s Import page.',
]:
    ws.cell(row=r, column=1, value=line).font = N; ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4); ws.cell(row=r, column=1).alignment = Alignment(wrap_text=True); ws.row_dimensions[r].height = 30; r += 1
r += 1
legend(ws, r)

# ------------------------------------------------------------------ 1 用户
ws = wb.create_sheet('1 用户 Users')
title(ws, '用户 Users', 'One row per person who will use the app', '每位使用者一行。用户名用 UBS 里的业务员代码（小写）。角色：salesperson 业务员 / finance 财务 / owner 老板。',
      'One row per user. Username = the salesperson code as in UBS (lower case). Roles: salesperson / finance / owner.')
cols = [('username 用户名', 16, 'Login name, lower case letters/digits, e.g. the UBS salesperson code', True),
        ('display_name 姓名', 20, 'Name shown in the app (Chinese is fine)', True),
        ('role 角色', 14, 'salesperson | finance | owner', True),
        ('salesperson_code 业务员代码', 20, 'Exactly as in UBS / seed_customers.csv (e.g. ALI). Leave empty for finance/owner.', True),
        ('password 初始密码', 16, 'Optional. If empty the app uses sk1234; every user must change it at first login.', True),
        ('phone 手机型号', 18, 'Not imported. Android/iPhone model, to check the app on that phone.', True),
        ('whatsapp 联系电话', 18, 'Not imported. For sending the link and password.', True),
        ('notes 备注', 30, 'Not imported.', True)]
header(ws, 5, cols)
r = rows(ws, 6, [('ali', '阿里 Ali', 'salesperson', 'ALI', '', 'Samsung A15', '012-xxxxxxx', '示例 example')], 8, {1, 2, 3, 4, 5, 6, 7, 8})
for c in range(1, 9): ws.cell(row=6, column=c).fill = GREEN
r = rows(ws, r, [], 8, {1, 2, 3, 4, 5, 6, 7, 8}, n_blank=14)
dv = DataValidation(type='list', formula1='"salesperson,finance,owner"', allow_blank=True); ws.add_data_validation(dv); dv.add(f'C6:C{r}')
legend(ws, r + 1)

# ------------------------------------------------------------------ 2 单位核对
ws = wb.create_sheet('2 单位核对 UOM')
title(ws, '单位核对（里程碑 1） UOM reconciliation (Milestone 1)', 'Purchase unit vs selling unit for every core SKU',
      '灰色列是系统从 UBS 记录推算的证据，只供参考。请填黄色列：采购单位、销售单位、换算系数（1个采购单位含多少销售单位）、核实成本（每销售单位）。没有采购记录的SKU，在 no_purchase_confirmed 填 Y。',
      'Grey columns are evidence from UBS records, for reference only. Fill the yellow columns: purchase unit, selling unit, factor (selling units per purchase unit), verified cost per selling unit. For SKUs with no purchases, put Y in no_purchase_confirmed.')
cols = [('item_code 货号', 12, 'Provided. Do not change.', False),
        ('description 品名', 26, 'Provided.', False),
        ('category 品类', 14, 'Provided.', False),
        ('uom_purchase_hint 采购记录单位', 14, 'Unit text as it appears in purchase records (evidence).', False),
        ('fy_purchase_qty 采购数量', 12, 'FY purchase quantity (evidence).', False),
        ('fy_purchase_value 采购金额', 14, 'FY purchase value RM (evidence).', False),
        ('implied_unit_cost 推算成本/采购单位', 14, 'Purchase value ÷ purchase qty = cost per PURCHASE unit (evidence).', False),
        ('uom_selling_hint 销售记录单位', 14, 'Unit text in sales records (evidence).', False),
        ('fy_sales_qty 销售数量', 12, 'FY sales qty (evidence).', False),
        ('implied_unit_price 推算售价/销售单位', 14, 'Sales value ÷ sales qty = price per SELLING unit (evidence).', False),
        ('uom_purchase 采购单位', 12, 'FILL: the unit we buy in, e.g. CTN, BALE, BAG, KG.', True),
        ('uom_selling 销售单位', 12, 'FILL: the unit we sell in, e.g. PKT, PCS, ROLL.', True),
        ('uom_factor 换算系数', 12, 'FILL: how many selling units in one purchase unit, e.g. 25 (25 packets per carton).', True),
        ('suggested_cost 推算成本/销售单位', 14, 'Formula: implied cost per purchase unit ÷ factor. A suggestion to check, not the answer.', False),
        ('verified_unit_cost 核实成本 (RM)', 16, 'FILL: the cost per SELLING unit you are confident in (check against a recent supplier invoice). Overrides everything once entered.', True),
        ('no_purchase_confirmed 确认无采购', 12, 'FILL: Y if this SKU genuinely had no purchases in the FY (e.g. old stock).', True),
        ('notes 备注', 30, 'Not imported. Questions, invoice reference, who checked.', True)]
header(ws, 5, cols)
ex = [('9.EC22', 'EC22A+LID', 'Disposables', 'CTN', 1720, 209496.00, 121.80, 'PKT', 41200, 4.86, 'CTN', 'PKT', 25, None, 4.87, '', '示例：25包/箱 → 121.80÷25 = 4.87  example'),
      ('70.JP9', 'JSP 9" plate', 'Disposables', 'BALE', 3400, 183464.00, 53.96, 'PKT', 9800, 19.31, '', '', None, None, None, '', '示例：待填 example, to fill'),
      ('SXLSK', 'Rubbish bag XL', 'Plastic bags', 'PKT', 58000, 129340.00, 2.23, 'PKT', 60100, 4.23, 'PKT', 'PKT', 1, None, 2.23, '', '示例：单位一致 example, units align')]
inputs = {11, 12, 13, 15, 16, 17}
r = rows(ws, 6, ex, 17, inputs)
for rr in range(6, 9):
    for c in range(1, 18): ws.cell(row=rr, column=c).fill = GREEN
r = rows(ws, r, [], 17, inputs, n_blank=180)
for rr in range(6, r):
    ws.cell(row=rr, column=14, value=f'=IF(AND(ISNUMBER(G{rr}),ISNUMBER(M{rr}),M{rr}>0),ROUND(G{rr}/M{rr},2),"")').font = N
    for c in (6, 7, 10, 14, 15): ws.cell(row=rr, column=c).number_format = '#,##0.00'
    for c in (5, 9): ws.cell(row=rr, column=c).number_format = '#,##0'
dv = DataValidation(type='list', formula1='"Y,N"', allow_blank=True); ws.add_data_validation(dv); dv.add(f'P6:P{r}')
ws.cell(row=r + 1, column=1, value='说明 Notes').font = H
ws.cell(row=r + 2, column=1, value='· 灰色列由应用的「导出 CSV」预填（产品与价格页面）；实际发出的表会列出全部 178 个核心 SKU，优先处理 cost_data_quality = UOM_SUSPECT 的 73 个。').font = S
ws.cell(row=r + 3, column=1, value='· Grey columns are pre-filled from the app\'s product export; the issued sheet lists all 178 core SKUs, with the 73 flagged UOM_SUSPECT first.').font = S
ws.cell(row=r + 4, column=1, value='· 示例数字来自 DATA_SPEC 第3节。 Example figures are from DATA_SPEC section 3.').font = S

# ------------------------------------------------------------------ 3 价格等级
ws = wb.create_sheet('3 价格等级 Tiers')
title(ws, '价格等级与规则 Price tiers and rules', 'Which customers get which price', '表A：等级名称（可增减）。表B：客户类型 × 规模 → 等级；* 代表任何。没有规则的组合用默认等级 STD。',
      'Table A: tier names (add or remove). Table B: customer type × size → tier; * means any. Combinations without a rule get the default tier STD.')
ws.cell(row=5, column=1, value='表A 价格等级 Table A · Tiers').font = H
cols = [('code 代码', 10, 'Short upper-case code, e.g. STD, KEY, SML', True), ('name_zh 中文名', 16, 'Name shown to salespeople', True), ('name_en English', 16, '', True), ('sort 排序', 8, 'Display order', True), ('who 适用客户 (说明)', 40, 'Not imported: who should get this tier', True)]
header(ws, 6, cols)
r = rows(ws, 7, [('STD', '标准价', 'Standard', 10, '一般客户 default'), ('KEY', '大客户价', 'Key account', 20, '例如 超市、大批发 e.g. hypermarkets, big wholesalers'), ('SML', '小客户价', 'Small account', 30, '例如 小杂货店 e.g. small kedai runcit')], 5, {1, 2, 3, 4, 5})
r = rows(ws, r, [], 5, {1, 2, 3, 4, 5}, n_blank=3)
r += 1
ws.cell(row=r, column=1, value='表B 等级规则 Table B · Tier rules (customer_type × size_tier → price_tier)').font = H; r += 1
cols = [('customer_type 客户类型', 22, 'Exactly as in seed_customers.csv, or * for any', True), ('size_tier 规模', 12, 'Exactly as in seed_customers.csv (e.g. S/M/L), or * for any', True), ('price_tier 等级', 12, 'A code from Table A', True), ('notes 备注', 40, '', True)]
header(ws, r, cols); hb = r
r = rows(ws, r + 1, [('Hypermarket', '*', 'KEY', '示例 example'), ('*', 'S', 'SML', '示例 example')], 4, {1, 2, 3, 4})
for rr in range(hb + 1, r):
    for c in range(1, 5): ws.cell(row=rr, column=c).fill = GREEN
r = rows(ws, r, [], 4, {1, 2, 3, 4}, n_blank=12)
ws.freeze_panes = None
legend(ws, r + 1)

# ------------------------------------------------------------------ 4 目标毛利
ws = wb.create_sheet('4 目标毛利 Margins')
title(ws, '目标毛利 Target margin by category', 'Used to derive list prices from verified cost', '目标毛利 = (售价 − 成本) ÷ 售价。底价折扣 = 允许低于目录价的最大百分比。这些数字通过「产品与价格 → 批量修改」应用到每个品类。',
      'Target margin = (price − cost) ÷ price. Floor discount = the largest discount off list a salesperson may give. Applied per category via Products → Bulk edit.')
cols = [('category 品类', 22, 'Exactly as in seed_products.csv', False), ('core_skus 核心SKU数', 12, 'Provided.', False), ('fy_sales_value 年销售额', 16, 'Provided.', False),
        ('target_margin_pct 目标毛利 %', 16, 'FILL: e.g. 25 means 25%', True), ('floor_discount_pct 底价折扣 %', 16, 'FILL: e.g. 5 means floor = list − 5%', True), ('notes 备注', 40, '', True)]
header(ws, 5, cols)
r = rows(ws, 6, [('Plastic bags', 40, 8500000, 18, 5, '示例 example')], 6, {4, 5, 6})
for c in range(1, 7): ws.cell(row=6, column=c).fill = GREEN
r = rows(ws, r, [], 6, {4, 5, 6}, n_blank=20)
for rr in range(6, r): ws.cell(row=rr, column=3).number_format = '#,##0'
ws.cell(row=r + 1, column=1, value='· 品类列表由应用导出预填。 The category list is pre-filled from the app export.').font = S
legend(ws, r + 3)

# ------------------------------------------------------------------ 5 核心价格
ws = wb.create_sheet('5 核心价格 Prices')
title(ws, '核心SKU价格表 Core SKU price list', 'Current list and floor price per tier', '每个等级填目录价和底价（每销售单位，RM）。没有底价就只填目录价（系统会把底价设为目录价）。生效日期留空 = 导入当天。',
      'For each tier, list and floor price per SELLING unit in RM. If there is no floor, fill list only (floor is set equal to list). Empty effective_from = the import date.')
cols = [('item_code 货号', 12, 'Provided.', False), ('description 品名', 26, 'Provided.', False), ('uom_selling 销售单位', 10, 'Provided after UOM reconciliation.', False),
        ('list_STD 标准目录价', 12, 'FILL: list price, STD tier', True), ('floor_STD 标准底价', 12, 'FILL: floor price, STD tier', True),
        ('list_KEY 大客户目录价', 12, 'FILL (optional)', True), ('floor_KEY 大客户底价', 12, 'FILL (optional)', True),
        ('list_SML 小客户目录价', 12, 'FILL (optional)', True), ('floor_SML 小客户底价', 12, 'FILL (optional)', True),
        ('effective_from 生效日期', 14, 'YYYY-MM-DD, optional', True), ('note 备注', 30, 'Imported into the price history note', True)]
header(ws, 5, cols)
r = rows(ws, 6, [('9.EC22', 'EC22A+LID', 'PKT', 6.10, 5.80, 5.95, 5.70, None, None, '2026-10-01', '示例 example')], 11, {4, 5, 6, 7, 8, 9, 10, 11})
for c in range(1, 12): ws.cell(row=6, column=c).fill = GREEN
r = rows(ws, r, [], 11, {4, 5, 6, 7, 8, 9, 10, 11}, n_blank=180)
for rr in range(6, r):
    for c in range(4, 10): ws.cell(row=rr, column=c).number_format = '#,##0.00'
ws.cell(row=r + 1, column=1, value='· 货号与品名由应用导出预填（178个核心SKU）。 Item codes are pre-filled from the app export (178 core SKUs).').font = S

# ------------------------------------------------------------------ 6 客户等级
ws = wb.create_sheet('6 客户等级 Customers')
title(ws, '客户等级 Customer tiers', 'Manual tier for key customers; confirm type, size and salesperson', '只需填写需要手动指定等级的客户（例如大客户）。price_tier_override 填 Y 表示手动指定，规则不会覆盖。其余客户按第3表的规则自动分配。',
      'Only fill customers that need a manual tier (e.g. key accounts). price_tier_override = Y means set by hand; rules will not overwrite it. Everyone else is assigned by the rules in sheet 3.')
cols = [('customer_code 客户代码', 14, 'As in UBS / seed_customers.csv', False), ('customer_name 客户名称', 26, 'Provided.', False), ('customer_type 客户类型', 16, 'Provided; correct if wrong', True),
        ('size_tier 规模', 10, 'Provided; correct if wrong', True), ('salesperson 业务员', 12, 'Provided; correct if wrong', True),
        ('price_tier 等级', 12, 'FILL: tier code from sheet 3', True), ('price_tier_override 手动', 10, 'Y = manual, rules will not overwrite', True), ('notes 备注', 36, '', True)]
header(ws, 5, cols)
r = rows(ws, 6, [('C0001', 'Hypermarket Pasir Mas 1', 'Hypermarket', 'L', 'ALI', 'KEY', 'Y', '示例 example')], 8, {3, 4, 5, 6, 7, 8})
for c in range(1, 9): ws.cell(row=6, column=c).fill = GREEN
r = rows(ws, r, [], 8, {3, 4, 5, 6, 7, 8}, n_blank=60)
dv = DataValidation(type='list', formula1='"Y,N"', allow_blank=True); ws.add_data_validation(dv); dv.add(f'G6:G{r}')
ws.cell(row=r + 1, column=1, value='· 客户列表由应用导出预填（771个客户，按年采购额排序）。 Customer list is pre-filled from the app export (771 customers, by annual value).').font = S

# ------------------------------------------------------------------ 7 竞争对手
ws = wb.create_sheet('7 竞争对手 Competitors')
title(ws, '竞争对手与市场参考 Competitors and market references', 'Who we lose to, and where branded prices can be checked', '表A：业务员常遇到的竞争对手（记录竞争价时可直接选）。表B：品牌货的公开价格来源（网址）。',
      'Table A: competitors salespeople meet most (offered as choices when capturing a price). Table B: public price sources for branded SKUs.')
ws.cell(row=5, column=1, value='表A 竞争对手 Table A · Competitors').font = H
cols = [('competitor_name 竞争对手名称', 26, 'Short name salespeople use', True), ('area 地区', 16, 'e.g. Kota Bharu, Pasir Mas', True), ('categories 主要品类', 30, 'What they compete on', True), ('notes 备注', 40, 'Strengths, typical discounting', True)]
header(ws, 6, cols)
r = rows(ws, 7, [('竞争对手A Competitor A', 'Kota Bharu', 'Plastic bags, disposables', '示例 example')], 4, {1, 2, 3, 4})
for c in range(1, 5): ws.cell(row=7, column=c).fill = GREEN
r = rows(ws, r, [], 4, {1, 2, 3, 4}, n_blank=12)
r += 1
ws.cell(row=r, column=1, value='表B 品牌货市场参考来源 Table B · Market-reference sources for branded SKUs').font = H; r += 1
cols = [('item_code 货号', 12, 'Branded SKU (is_branded = 1)', True), ('brand 品牌', 14, 'e.g. Seamaster, LYG, Nestlé', True), ('source_name 来源名称', 22, 'Wholesaler / marketplace name', True), ('source_url 网址', 50, 'Full https:// link to the product page', True), ('unit 单位', 10, 'e.g. CTN 24', True), ('notes 备注', 30, '', True)]
header(ws, r, cols); hb = r
r = rows(ws, r + 1, [('CO043', 'Seamaster', 'Wholesaler X', 'https://example.com/seamaster-250ml-x24', 'CTN 24', '示例 example')], 6, {1, 2, 3, 4, 5, 6})
for c in range(1, 7): ws.cell(row=hb + 1, column=c).fill = GREEN
r = rows(ws, r, [], 6, {1, 2, 3, 4, 5, 6}, n_blank=20)
ws.freeze_panes = None
ws.cell(row=r + 1, column=1, value='· 参考价永远只供参考，不会自动调价。 Reference prices are for reference only; nothing is applied automatically.').font = S

# ------------------------------------------------------------------ 8 决策
ws = wb.create_sheet('8 决策 Decisions')
title(ws, '老板决策清单 Owner decisions', 'Rules the app enforces; the owner decides them', '每项写下决定和负责人。默认值是应用现在的设定。', 'Write the decision and the person responsible for each item. Defaults are what the app does today.')
cols = [('# ', 4, '', False), ('问题 Question', 56, '', False), ('选项 / 默认 Options / default', 40, '', False), ('决定 Decision', 30, 'FILL', True), ('负责人 Owner', 14, 'FILL', True), ('备注 Notes', 30, '', True)]
header(ws, 5, cols)
qs = [
    ('业务员能否看到成本和毛利？\nMay salespeople see cost and margin?', '默认：不能（只看目录价和底价）。可随时在设置中开启。\nDefault: no (list and floor only). Can be switched on in Settings.'),
    ('谁可以修改成本、目标毛利和价格？\nWho may change cost, target margin and prices?', '默认：财务 (finance) 和老板 (owner)。\nDefault: finance and owner.'),
    ('底价的含义：业务员未经批准可给的最低价？\nMeaning of floor price: lowest a salesperson may quote without approval?', '默认：是。低于底价不阻止，但会被记录；v1 没有审批流程。\nDefault: yes. Going below is not blocked but is recorded; no approval workflow in v1.'),
    ('默认价格等级 Default price tier', '默认：STD\nDefault: STD'),
    ('默认底价折扣（目录价减多少 % = 底价）\nDefault floor discount (list − x % = floor)', '默认：5%\nDefault: 5%'),
    ('OK 质量但未核实成本的SKU，是否允许显示推算毛利（标注「未核实」）？\nFor OK-quality SKUs without a verified cost, may the app show the implied margin labelled "unverified"?', '默认：显示并标注。更严格的做法：一律不显示。\nDefault: shown with the label. Stricter option: never shown.'),
    ('先从178个核心SKU开始，其余915个SKU暂不定价？\nStart with the 178 core SKUs only; leave the other 915 unpriced for now?', '默认：是（规格建议）。\nDefault: yes (as the brief recommends).'),
    ('客户类型和规模的口径由谁确认？\nWho confirms the customer type and size classification?', '建议：老板 + 各业务员核对自己的客户。\nSuggested: owner + each salesperson checks their own customers.'),
    ('销售 / 采购数据多久重新导入一次？\nHow often will UBS sales/purchase exports be re-imported?', '建议：每月一次，由文员在「导入」页面上传。\nSuggested: monthly, by the clerk on the Import page.'),
    ('业务员手机：安卓/苹果？是否只用中文界面？\nSalesperson phones: Android/iPhone? Chinese-only UI?', '应用是中文为主、英文为辅；请确认是否需要简体或繁体。\nUI is Chinese-first with English; confirm Simplified vs Traditional.'),
    ('上线日期和培训时间 Go-live date and training session', ''),
    ('品牌货市场参考价：是否需要定期查询（里程碑7）？\nBranded market reference: do you want scheduled lookups (Milestone 7)?', '默认：先手动录入，M1–M6 用起来后再说。\nDefault: manual entry first; revisit after M1–M6 are in use.'),
]
r = 6
for i, (q, o) in enumerate(qs, 1):
    vals = [i, q, o, None, None, None]
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=r, column=c, value=v); cell.border = BORDER; cell.alignment = Alignment(wrap_text=True, vertical='top')
        cell.font = INPUT if c >= 4 else N
        if c >= 4: cell.fill = YELLOW
    ws.row_dimensions[r].height = 60; r += 1
legend(ws, r + 1)

for sheet in wb.worksheets:
    for row in sheet.iter_rows():
        for cell in row:
            if cell.font and cell.font.name != FONT:
                cell.font = Font(name=FONT, size=cell.font.size or 10, bold=cell.font.bold, color=cell.font.color)
wb.save('SK_Keong_Pricing_App_Data_Request.xlsx')
print('saved')
